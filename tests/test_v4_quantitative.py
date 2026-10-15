"""Q1 numeric/storage regressions; synthetic fixtures never prove source accuracy."""
import json
from pathlib import Path
import pytest
from bct.v4 import quantitative as q
from bct.v4.store import Ledger,sha,file_sha


def meta(tmp_path,raw,version='one'):
    p=tmp_path/(version+'.raw');p.write_bytes(raw)
    return {'source_id':'synthetic','url':'https://example.org/test','sha256':sha(raw),'path':str(p),'published_at':'2026-09-18','acquired_at':'2026-10-15T00:00:00+00:00'}


@pytest.mark.parametrize('v,number,status',[(0,'0','REPORTED'),('-1.5','-1.5','REPORTED'),('1,234','1234','REPORTED'),('e123','123','ESTIMATE'),('(52)','-52','REPORTED'),('W',None,'WITHHELD'),('D',None,'WITHHELD'),('—',None,'MISSING'),('N/A',None,'MISSING'),('<0.5',None,'UNKNOWN_FORMAT'),('1 million',None,'UNKNOWN_FORMAT'),('inf',None,'UNKNOWN_FORMAT')])
def test_exact_numeric_missing_not_zero(v,number,status):assert q.number(v)==(number,status)


def row(m,value='100',unit='USD',kind='MONEY',month='2026-01'):
    return q.observation(m,'case','product','test',value,unit,kind,month,{'format':'JSON','path':['value']})


def test_money_scale_is_not_a_product_count(tmp_path):
    m=meta(tmp_path,b'{}');r=q.observation(m,'case','product','investment','800','million USD','MONEY','2026-01',{},scale='1000000')
    assert r['value_normalized_or_null']=='800000000'and r['measurement_type']=='MONEY'
    assert r['stage']=='UNKNOWN'and not r['automatic_promotion']


def test_monthly_saar_is_not_monthly_unit_count(tmp_path):
    raw=b'"TOTASS: Vehicles"\n"TOTASS" 2026 1.1 2.2\n';m=meta(tmp_path,raw)
    out=q.parse_fed(raw,m,['TOTASS'],'million vehicles/year SAAR','PHYSICAL','assemblies','2026-01','2026-02')
    assert len(out)==2 and out[1]['period_start']=='2026-02-01'
    assert out[1]['value_normalized_or_null']=='2200000.0'
    assert out[1]['stock_or_flow']=='ANNUALIZED_FLOW_RATE'
    assert q.verify_locator(out[1],raw)


def test_same_input_noop_restart_append_revision_unit_isolation(tmp_path):
    l=Ledger(tmp_path/'ledger');m=meta(tmp_path,b'{"value":100}')
    r=row(m);assert q.commit_page(l,'fixed',m,[r])['added']==1
    before=file_sha(l.path);assert q.commit_page(l,'fixed',m,[r])['status']=='UNCHANGED';assert file_sha(l.path)==before
    l.close();l=Ledger(tmp_path/'ledger');assert q.cursor(l,'fixed')['next']is None
    m2=meta(tmp_path,b'{"value":110}','two');r2=row(m2,'110')
    result=q.commit_page(l,'fixed',m2,[r2]);assert result['changes']=={'REVISED_SAME_PERIOD':1}
    m3=meta(tmp_path,b'{"value":120}','three');r3=row(m3,'120',unit='thousand USD')
    assert q.commit_page(l,'fixed',m3,[r3])['changes']=={'UNIT_CHANGED_QUARANTINED':1}
    assert q.records(l)[-1]['record']['quality_state']=='HELD_UNIT_CHANGE'
    m4=meta(tmp_path,b'{"value":130}','four');r4=row(m4,'130',month='2026-02')
    assert q.commit_page(l,'fixed',m4,[r4])['changes']=={'NEW_OBSERVATION':1}
    assert len(q.records(l))==4;l.close()


def test_gap_and_crash_atomic_resume(tmp_path):
    l=Ledger(tmp_path/'ledger');m=meta(tmp_path,b'{}');r=row(m)
    q.commit_page(l,'paged',m,[r],page=2,has_next=False)
    assert q.cursor(l,'paged')=={'next':1,'committed_pages':1,'gap':True}
    with pytest.raises(ValueError,match='DUPLICATE_ROW'):
        q.commit_page(l,'paged',m,[r,r],page=1,has_next=True)
    assert q.cursor(l,'paged')['next']==1
    q.commit_page(l,'paged',m,[r],page=1,has_next=True);assert q.cursor(l,'paged')['next']is None
    l.close()


def test_hash_and_source_location_tampering(tmp_path):
    m=meta(tmp_path,b'{"value":100}');r=row(m)
    assert q.verify_locator(r,Path(m['path']).read_bytes())
    with pytest.raises(ValueError,match='HASH'):q.verify_locator(r,b'{"value":101}')
    r['source_locator']['path']=['absent']
    with pytest.raises(KeyError):q.verify_locator(r,Path(m['path']).read_bytes())
    l=Ledger(tmp_path/'l');Path(m['path']).write_bytes(b'changed')
    with pytest.raises(ValueError,match='HASH'):q.commit_page(l,'source',m,[r])
    assert not q.records(l);l.close()


def test_size_rejection_rolls_back_observations_and_cursor(tmp_path,monkeypatch):
    l=Ledger(tmp_path/'l');m=meta(tmp_path,b'{}');monkeypatch.setattr(q,'LIMIT',1)
    with pytest.raises(ValueError,match='COMPRESSED_LIMIT'):q.commit_page(l,'source',m,[row(m)])
    assert q.cursor(l,'source')['next']==1 and q.records(l)==[];l.close()


def test_eight_table_restore_is_local_not_remote(tmp_path):
    l=Ledger(tmp_path/'l');m=meta(tmp_path,b'{}');q.commit_page(l,'source',m,[row(m)])
    result=q.backup_restore(l,tmp_path/'backup/v4.sqlite3')
    assert result['observations']==1 and result['not_remote_private_store']
    restored=Ledger(tmp_path/'backup');assert q.records(restored)==q.records(l);restored.close();l.close()


def test_ted_single_total_array_without_inventing_lot_relationship(tmp_path):
    raw=json.dumps({'notices':[{'publication-number':'synthetic-1','publication-date':'2026-09-16+02:00','total-value':[100],'total-value-cur':['EUR'],'quantity-lot':[2,10],'quantity-unit-lot':['kg','tonnes']}],'totalNoticeCount':1,'timedOut':False}).encode();m=meta(tmp_path,raw)
    rows,p,held=q.parse_ted(raw,m);assert len(rows)==1 and len(held)==2
    assert rows[0]['measurement_type']=='MONEY'and not rows[0]['attributes']['is_actual_order']
    assert q.verify_locator(rows[0],raw)
    p['notices'][0]['total-value']=[100,200]
    assert q.parse_ted(json.dumps(p).encode(),m)[0]==[]


def test_no_source_bodies_or_s_gate_writes(tmp_path):
    l=Ledger(tmp_path/'l');m=meta(tmp_path,b'{}');q.commit_page(l,'source',m,[row(m)])
    for table in ['documents','blobs','readings','runs','firsts']:assert l.db.execute('select count(*)from '+table).fetchone()==(0,)
    assert q.records(l)[0]['record']['backtest_state'].startswith('BACKTEST_UNAVAILABLE');l.close()


def test_usgs_footnote_is_not_negative_and_withheld_is_not_zero(tmp_path):
    import openpyxl
    from io import BytesIO
    w=openpyxl.Workbook();s=w.active;s.title='T1';s['A2']='Data are in metric tons, tungsten content.'
    s['B4']='Consumption';s['C4']='Imports for consumption';s['D4']='Stocks, end of period';s['E4']='Net production';s['F4']='Stocks'
    s['A5']='2026';s['A6']='January'
    for cell,val in [('B6','W'),('C6',0),('D6','(2)'),('E6',100),('F6','D')]:s[cell]=val
    buf=BytesIO();w.save(buf);raw=buf.getvalue();m=meta(tmp_path,raw)
    rows,holds=q.parse_tungsten(raw,m,'2026-01','2026-01');assert len(rows)==5
    assert [r['value_decimal_or_null']for r in rows]==[None,'0',None,'100',None]
    assert rows[2]['unknown_reason']=='FOOTNOTE_REFERENCE_NOT_NEGATIVE_QUANTITY'
    assert rows[1]['source_lineage_id']=='US_CENSUS_INTLTRADE'
    assert q.verify_locator(rows[2],raw)


def test_xlsx_m3_month_columns_scale_and_same_source_location(tmp_path):
    import openpyxl
    from io import BytesIO
    w=openpyxl.Workbook();s=w.active;s.title='synthetic-table';s['A1']='AMTMNO';s['B1']=2026
    for i in range(3,15):s.cell(1,i,i*10)
    buf=BytesIO();w.save(buf);raw=buf.getvalue();m=meta(tmp_path,raw)
    rows=q.parse_m3(raw,m,['AMTM'],'orders','2026-02','2026-03')
    assert [r['source_locator']['cell']for r in rows]==['D1','E1']
    assert [r['value_normalized_or_null']for r in rows]==['40000000','50000000']
    assert all(q.verify_locator(r,raw)for r in rows)


def test_json_decimal_cents_never_round_through_binary_float(tmp_path):
    raw=b'{"notices":[{"publication-number":"synthetic","publication-date":"2026-09-16","total-value":9007199254740992.13,"total-value-cur":"EUR"}],"totalNoticeCount":1}'
    m=meta(tmp_path,raw);rows,_,_=q.parse_ted(raw,m)
    assert rows[0]['value_decimal_or_null']=='9007199254740992.13'
    assert q.verify_locator(rows[0],raw)


def test_secret_url_refused_and_local_state_conflict_does_not_overwrite(tmp_path):
    m=meta(tmp_path,b'{}');m['url']='https://example.org/data?api_key=SYNTHETIC_NOT_A_KEY'
    with pytest.raises(ValueError,match='SECRET_BEARING'):row(m)
    m['url']='https://example.org/data';l=Ledger(tmp_path/'l');q.commit_page(l,'one',m,[row(m)])
    dest=tmp_path/'backup/v4.sqlite3';before=file_sha(l.path)
    with pytest.raises(ValueError,match='CAS_HASH_CHANGED'):q.backup_restore(l,dest,expected_current_hash='0'*64)
    assert not dest.exists()and file_sha(l.path)==before
    assert q.backup_restore(l,dest,expected_current_hash=before)['observations']==1;l.close()


def test_validation_history_controls_handoff_and_objective_binding(tmp_path):
    l=Ledger(tmp_path/'l');m=meta(tmp_path,b'{}');q.commit_page(l,'one',m,[row(m)])
    item=q.records(l)[0]
    with l.db:l.event('QUANTITATIVE_VALIDATION',{'record_ids':[item['id']],'quality_state':'HELD_PERIOD_UNKNOWN','reason':'Contract start is not measurement period.'})
    result=q.research_export(l)
    assert result['objective_version']=='bct-objective-1'
    assert result['observations'][0]['record']['quality_state']=='HELD_PERIOD_UNKNOWN'
    assert q.records(l)[0]['record']['quality_state']=='SOURCE_LOCATED_PENDING_REVIEW'
    assert result['candidates']==[];l.close()


def test_contract_period_correction_view_preserves_original_draft(tmp_path):
    l=Ledger(tmp_path/'l');m=meta(tmp_path,b'{}');q.commit_page(l,'one',m,[row(m)])
    item=q.records(l)[0]
    with l.db:l.event('QUANTITATIVE_VALIDATION',{'record_ids':[item['id']],'quality_state':'HELD_OBSERVATION_PERIOD_UNKNOWN','reason':'contract dates are not this statistic period'})
    out=q.research_export(l)['observations'][0]['record']
    assert out['period_start']=='UNKNOWN'and out['period_end']=='UNKNOWN'
    assert q.records(l)[0]['record']['period_start']=='2026-01-01';l.close()


def source_tool(tmp_path,monkeypatch,raw,received):
    import importlib.util
    spec=importlib.util.spec_from_file_location('q1_source_tool',Path(__file__).resolve().parents[1]/'tools/v4_quantitative_sources.py')
    tool=importlib.util.module_from_spec(spec);spec.loader.exec_module(tool)
    monkeypatch.chdir(tmp_path);p=tmp_path/'docs/v4/quantitative';p.mkdir(parents=True)
    (p/'source-registry.json').write_text(json.dumps({'file_tasks':[{'key':'m3-orders','source_id':'m3','url':'https://example.org/official','sha256':sha(raw),'acquired_at':'2026-10-01'}]}))
    monkeypatch.setattr(tool,'ROOT',tmp_path/'native');calls=[]
    class ExistingClient:
        def fetch(self,url):calls.append(url);return received,url,200,'application/octet-stream'
    monkeypatch.setattr(tool,'SourceClient',ExistingClient)
    return tool,calls


def test_frozen_source_different_version_held_without_repeated_download(tmp_path,monkeypatch):
    tool,calls=source_tool(tmp_path,monkeypatch,b'frozen source',b'newer version')
    result=tool.prepare('m3');assert result['results'][0]['status']=='NEW_VERSION_HELD_NOT_REPLACED'
    assert not (tool.ROOT/'m3-orders.raw').exists()
    assert tool.prepare('m3')['requests']==0 and len(calls)==1
    assert list((tool.ROOT/'unfrozen-versions').glob('*.raw'))


def test_exact_source_restore_then_noop_preserves_original_acquisition(tmp_path,monkeypatch):
    tool,calls=source_tool(tmp_path,monkeypatch,b'exact source',b'exact source')
    assert tool.prepare('m3')['results'][0]['status']=='EXACT_VERSION_RESTORED'
    assert json.loads((tool.ROOT/'m3-orders.json').read_text())['acquired_at']=='2026-10-01'
    assert tool.prepare('m3')['requests']==0 and len(calls)==1


def test_cached_checkpoint_budget_counts_real_gzip_members_and_mutation():
    import gzip
    raw=b'abc'*50000+bytes(range(256))*900
    budget=q.CheckpointBudget()
    def packed(value):
        return b''.join(gzip.compress(value[i:i+65536],compresslevel=1,mtime=0)for i in range(0,len(value),65536))
    assert budget.size(raw)==len(packed(raw))
    assert budget.size(raw)==len(packed(raw))
    changed=raw[:1234]+b'changed'+raw[1241:]
    assert budget.size(changed)==len(packed(changed))
    assert gzip.decompress(packed(changed))==changed


def test_official_hs_leaf_units_not_parent_totals_or_tungsten_content_conversion(tmp_path):
    raw=json.dumps([{'htsno':'8101','description':'synthetic root','units':[]},
      {'htsno':'8101.10.00.00','description':'synthetic powder','units':['kg','W kg']},
      {'htsno':'8102.10.00.00','description':'other material','units':['kg']}]).encode()
    m=meta(tmp_path,raw);mapped=q.parse_hs_mapping(raw,m,'8101')
    assert len(mapped)==1 and mapped[0]['import_HTS10']=='8101100000'
    assert mapped[0]['primary_quantity_unit']=='kg'and mapped[0]['secondary_quantity_unit']=='W kg'
    assert mapped[0]['trade_quantities']=='NOT_FETCHED'
    with pytest.raises(ValueError,match='ROOT_NOT_IN_OFFICIAL'):q.parse_hs_mapping(raw,m,'2504')


def test_report_roles_preserve_plan_amount_and_broad_production_distinctions(tmp_path):
    m=meta(tmp_path,b'{}');r=row(m)
    assert q.quantity_role(r)=='MONEY_ONLY_NOT_PHYSICAL_DEMAND'
    r.update(source_id='eia',metric='nameplate_capacity',measurement_type='PHYSICAL',attributes={'sheet':'Planned'})
    assert q.quantity_role(r)=='PLANNED_NAMEPLATE_NOT_AVAILABLE'
    r['attributes']['sheet']='Operating'
    assert q.quantity_role(r)=='OPERATING_NAMEPLATE_NOT_QUALIFIED'
    r.update(source_id='g17',metric='vehicle_assemblies')
    assert q.quantity_role(r)=='REPORTED_PRODUCTION_ANNUALIZED_RATE'
    r['value_decimal_or_null']=None
    assert q.quantity_role(r)=='WITHHELD_OR_UNKNOWN'
