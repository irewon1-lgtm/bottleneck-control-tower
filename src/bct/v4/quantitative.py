"""Q1 numeric observations in the existing append-only V4 research Ledger.

No candidate creation, CPU/AI execution, schedule changes or financial-to-physical
conversion. Raw official files remain in an isolated private research cache.
"""
from calendar import monthrange
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import gzip
from pathlib import Path
from urllib.parse import urlsplit,parse_qsl
import re
import sqlite3
from .store import Ledger, encode, now, sha, file_sha

PARSER = 'bct-q1-numeric-1'
OBS = 'QUANTITATIVE_OBSERVATION'
PAGE = 'QUANTITATIVE_PAGE_COMMITTED'
LIMIT = 16 * 1024 * 1024
MISSING = {'', '-', '--', '—', 'NA', 'N/A', 'N.A.', '(NA)', 'W', '(W)', 'D', '(D)', 'S', '(S)', 'X', '..', '...'}


def number(value):
    """Keep absent/suppressed/estimate flags; zero and negative remain numbers."""
    text = '' if value is None else str(value).strip()
    if text.upper() in MISSING:
        return None, ('WITHHELD' if text.upper().strip('()') in {'W','D','S'} else 'MISSING')
    flag = 'ESTIMATE' if re.fullmatch(r'e\s*[-+\d,.]+', text, re.I) else 'REPORTED'
    cleaned = re.sub(r'^e\s*', '', text, flags=re.I) if flag == 'ESTIMATE' else text
    if cleaned.endswith('e') and re.fullmatch(r'[-+\d,.]+e',cleaned,re.I):
        flag = 'ESTIMATE';cleaned=cleaned[:-1]
    if re.fullmatch(r'\(\d[\d,.]*\)', cleaned): cleaned='-'+cleaned[1:-1]
    if not re.fullmatch(r'[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|[-+]?\.\d+', cleaned):
        return None, 'UNKNOWN_FORMAT'
    try:
        v=Decimal(cleaned.replace(',',''))
        return (format(v,'f'), flag) if v.is_finite() else (None,'UNKNOWN_FORMAT')
    except InvalidOperation:return None,'UNKNOWN_FORMAT'


def period(month):
    if not re.fullmatch(r'\d{4}-\d{2}', month):raise ValueError('INVALID_MONTH')
    y,m=map(int,month.split('-'));return month+'-01',month+'-'+str(monthrange(y,m)[1])


def observation(meta, entity, classification, metric, value, unit, kind, month, locator,
                *, scale='1', adjustment='NOT_APPLICABLE', stock_or_flow='UNKNOWN',
                status=None, geography='US', attributes=None, lineage=None):
    if any(re.search(r'key|token|auth|secret|password|session',k,re.I)for k,v in parse_qsl(urlsplit(meta['url']).query)):
        raise ValueError('SECRET_BEARING_SOURCE_URL')
    val,flag=number(value);start,end=period(month) if re.fullmatch(r'\d{4}-\d{2}',month)else(month,month)
    if kind not in {'PHYSICAL','MONEY','INDEX','RATIO'}:raise ValueError('INVALID_MEASUREMENT_TYPE')
    normalized=None if val is None else format(Decimal(val)*Decimal(scale),'f')
    quality='SOURCE_LOCATED_PENDING_REVIEW' if val is not None else 'HELD'
    pub=meta.get('published_at','UNKNOWN')
    return {'source_id':meta['source_id'],'source_document_id':meta['sha256'],
      'source_lineage_id':lineage or meta.get('source_lineage_id',meta['source_id']),
      'source_url':meta['url'],'raw_or_safe_payload_hash':meta['sha256'],
      'original_version':meta.get('original_version',meta['sha256']), 'parser_version':PARSER,
      'source_locator':locator,'entity_or_facility_id':str(entity),
      'product_or_classification':str(classification),'sector_hint':meta.get('sector_hint','UNCLASSIFIED'),
      'geography':geography,'metric':metric,'value_original':None if value is None else str(value),
      'value_decimal_or_null':val,'value_normalized_or_null':normalized,
      'unit_original':unit,'unit_normalized':{'million USD':'USD','million vehicles/year SAAR':'vehicles/year SAAR'}.get(unit,unit),
      'scale':scale,'measurement_type':kind,'stock_or_flow':stock_or_flow,'adjustment':adjustment,
      'period_start':start,'period_end':end,'published_at':pub,'acquired_at':meta['acquired_at'],
      'revised_at':meta.get('last_modified','UNKNOWN'),'vintage':meta.get('vintage',pub),
      'status':status or flag,'coverage_scope':meta.get('coverage_scope','MANIFEST_ONLY'),
      'rights_scope':meta.get('rights_scope','PRIVATE_OFFICIAL_STATISTICAL_RESEARCH'),
      'original_available_for_replay':True,'quality_state':quality,
      'unknown_reason':flag if val is None else None,
      'backtest_state':'BACKTEST_UNAVAILABLE_NO_HISTORICAL_VINTAGE' if not meta.get('historical_vintage_verified') else 'VINTAGE_VERIFIED',
      'semantic_validation':'STATISTICAL_OBSERVATION_NOT_QUALIFIED_SUPPLY_OR_ORDER',
      'attributes':attributes or {},'stage':'UNKNOWN','automatic_promotion':False}


def obs_key(r):
    return sha([r[k] for k in ('source_id','entity_or_facility_id','product_or_classification','geography','metric','period_start','period_end','adjustment','stock_or_flow')])


class CheckpointBudget:
    """Exact size of a valid concatenated-gzip checkpoint, cached by64KiB block.

    Standard gzip supports multiple members. Changed SQLite pages only cause
    their members to be recompressed; every block hash is still checked each
    transaction. No raise of the existing16MiB bound and no skipped size check.
    """
    def __init__(self):self.blocks={}
    def size(self, raw):
        total=0;current={}
        for at in range(0,len(raw),65536):
            block=raw[at:at+65536];digest=sha(block)
            length=self.blocks.get(digest)
            if length is None:length=len(gzip.compress(block,compresslevel=1,mtime=0))
            total+=length;current[digest]=length
        self.blocks=current
        return total


def commit_page(ledger, task, meta, rows, *, page=1, has_next=False, total=None):
    """Atomic page + output checkpoint; exact repeat is a no-op, versions append."""
    if page<1 or type(has_next)is not bool:raise ValueError('INVALID_PAGE_METADATA')
    if file_sha(meta['path']) != meta['sha256']:raise ValueError('SOURCE_HASH_MISMATCH')
    pid=sha([task,page,meta['sha256'],PARSER])
    # JSON expression indexes retain the existing event store and avoid rereading
    # every historical source record for every small resumed page.
    with ledger.db:
        ledger.db.execute("CREATE INDEX IF NOT EXISTS q1_page_id ON events(json_extract(payload,'$.id')) WHERE kind='QUANTITATIVE_PAGE_COMMITTED'")
        ledger.db.execute("CREATE INDEX IF NOT EXISTS q1_task_page ON events(json_extract(payload,'$.task'),json_extract(payload,'$.page')) WHERE kind='QUANTITATIVE_PAGE_COMMITTED'")
        ledger.db.execute("CREATE INDEX IF NOT EXISTS q1_observation_key ON events(json_extract(payload,'$.key')) WHERE kind='QUANTITATIVE_OBSERVATION'")
    if ledger.db.execute("select 1 from events where kind='QUANTITATIVE_PAGE_COMMITTED' and json_extract(payload,'$.id')=?",(pid,)).fetchone():
        return {'status':'UNCHANGED','added':0,'page':page}
    previous={};keys=list({obs_key(r)for r in rows})
    for start in range(0,len(keys),200):
        block=keys[start:start+200];marks=','.join('?'for _ in block)
        for payload, in ledger.db.execute("select payload from events where kind='QUANTITATIVE_OBSERVATION' and json_extract(payload,'$.key') in ("+marks+") order by seq",block):
            r=json.loads(payload);previous[r['key']]=r
    seen=set();added=0;changes=[]
    with ledger.db:
        for row in rows:
            required=('source_document_id','source_locator','value_original','unit_original','period_start','published_at','measurement_type')
            if any(k not in row for k in required):raise ValueError('INCOMPLETE_NUMERIC_RECORD')
            if row['raw_or_safe_payload_hash']!=meta['sha256']:raise ValueError('RECORD_HASH_MISMATCH')
            key=obs_key(row)
            if key in seen:raise ValueError('DUPLICATE_ROW_IDENTITY_WITHIN_PAGE')
            seen.add(key);old=previous.get(key);change='NEW_OBSERVATION'
            if old:
                before=old['record']
                if before['source_document_id']==row['source_document_id'] and before['source_locator']==row['source_locator'] and before['parser_version']==row['parser_version']:continue
                if (before['unit_original'],before['scale'])!=(row['unit_original'],row['scale']):change='UNIT_CHANGED_QUARANTINED';row=dict(row,quality_state='HELD_UNIT_CHANGE')
                elif before['value_decimal_or_null']!=row['value_decimal_or_null']:change='REVISED_SAME_PERIOD'
                else:change='NEW_VERSION_SAME_VALUE'
            ledger.event(OBS,{'key':key,'id':sha(row),'task':task,'change':change,'previous_id':old['id']if old else None,'record':row});added+=1;changes.append(change)
        ledger.event(PAGE,{'id':pid,'task':task,'page':page,'has_next':has_next,'total':total,'raw_hash':meta['sha256'],'parser':PARSER,'rows':len(rows),'added':added,'committed_at':now()})
        # Existing16MiB bound applies to the gzip checkpoint, not SQLite bytes.
        # Test inside the transaction so an oversized page rolls back entirely.
        budget=getattr(ledger,'_q1_checkpoint_budget',None)
        if budget is None:budget=ledger._q1_checkpoint_budget=CheckpointBudget()
        if budget.size(ledger.db.serialize())>LIMIT:
            raise ValueError('CHECKPOINT_16MIB_COMPRESSED_LIMIT')
    return {'status':'COMMITTED','added':added,'page':page,'changes':{c:changes.count(c)for c in set(changes)}}


def cursor(ledger, task):
    pages={}
    for p, in ledger.db.execute("select payload from events where kind='QUANTITATIVE_PAGE_COMMITTED' and json_extract(payload,'$.task')=? order by seq",(task,)):
        r=json.loads(p)
        if r['task']==task:pages[r['page']]=r
    n=1
    while n in pages:
        if not pages[n]['has_next']:return {'next':None,'committed_pages':len(pages),'gap':False}
        n+=1
    return {'next':n,'committed_pages':len(pages),'gap':any(k>n for k in pages)}


def records(ledger):
    return [json.loads(x[0])for x in ledger.db.execute('select payload from events where kind=? order by seq',(OBS,))]


def parse_fed(raw,meta,series,unit,kind,metric,start,end):
    rows=[];names={}
    for lineno,line in enumerate(raw.decode('utf-8-sig').splitlines(),1):
        title=re.fullmatch(r'"([^":]+): (.*)"',line.strip())
        if title:names[title[1]]=title[2];continue
        m=re.match(r'^"([^"\n]+)"\s+(\d{4})\s+(.*)$',line)
        if not m or m[1]not in series:continue
        tokens=m[3].split()
        if not 1<=len(tokens)<=12:raise ValueError('FED_INVALID_MONTH_COLUMNS')
        for month,v in enumerate(tokens,1):
            ym=m[2]+'-%02d'%month
            if start<=ym<=end:
                rows.append(observation(meta,m[1],names.get(m[1],m[1]),metric,v,unit,kind,ym,
                  {'format':'FED_TEXT','line':lineno,'token':month+1,'series_token':m[1],'year_token':m[2],'month':month},
                  scale='1000000'if kind=='PHYSICAL'else'1',adjustment='SAAR'if kind=='PHYSICAL'else'SA',stock_or_flow='ANNUALIZED_FLOW_RATE'if kind=='PHYSICAL'else'INDEX_LEVEL'))
    return rows


def parse_m3(raw,meta,prefixes,metric,start,end):
    import openpyxl
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True);out=[]
    for sheet in w:
        for row in sheet.iter_rows():
            code,year=row[0].value,row[1].value
            if not isinstance(code,str)or not any(code.startswith(p)for p in prefixes)or not isinstance(year,(int,float)):continue
            for month in range(1,13):
                ym='%04d-%02d'%(year,month)
                if not start<=ym<=end:continue
                cell=row[month+1]
                out.append(observation(meta,code,code,metric,cell.value,'million USD','MONEY',ym,
                  {'format':'XLSX','sheet':sheet.title,'cell':cell.coordinate,'series_cell':'A'+str(cell.row),'year_cell':'B'+str(cell.row),'month_column':month},scale='1000000',adjustment='SA',stock_or_flow='STOCK_MONTH_END'if metric in {'inventory','unfilled'}else'MONTHLY_FLOW'))
    w.close();return out


def parse_eia(raw,meta,ym,states):
    from .structured_changes import snapshot
    s=snapshot(raw,ym,meta['url']);out=[]
    if s['duplicate_ids_quarantined']:raise ValueError('EIA_DUPLICATE_FACILITY_KEYS')
    for key,r in s['records'].items():
        f=r['fields']
        if f.get('Plant State')not in states:continue
        metric='Nameplate Capacity (MW)'
        if metric not in f:continue
        attributes={k:f.get(k,'UNKNOWN')for k in ['Status','Technology','Operating Month','Operating Year','Planned Operation Month','Planned Operation Year']}
        attributes.update(sheet=r['sheet'],qualified_supply='UNKNOWN',actuality='PLAN'if r['sheet']=='Planned'else'REPORTED_STATUS_NOT_DELIVERY')
        out.append(observation(meta,'/'.join(key),f.get('Technology','UNKNOWN'),'nameplate_capacity',f[metric],'MW','PHYSICAL',ym,
          {'format':'XLSX','sheet':r['sheet'],'cell':r['cells'][metric]},geography=f['Plant State'],stock_or_flow='REPORTED_NAMEPLATE_STOCK',status='PLAN'if r['sheet']=='Planned'else'REPORTED',attributes=attributes))
    return out,s


def verify_locator(record,raw):
    if sha(raw)!=record['raw_or_safe_payload_hash']:raise ValueError('SOURCE_HASH_MISMATCH')
    loc=record['source_locator'];v=None
    if loc['format']=='SEC_INLINE':
        v=raw.decode('utf-8')[loc['start']:loc['end']]
    elif loc['format']=='FED_TEXT':
        import shlex
        tokens=shlex.split(raw.decode('utf-8-sig').splitlines()[loc['line']-1]);v=tokens[loc['token']]
        if tokens[0]!=loc['series_token']or tokens[1]!=loc['year_token']:raise ValueError('FED_IDENTITY_MISMATCH')
    elif loc['format']=='XLSX':
        import openpyxl
        w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True);v=w[loc['sheet']][loc['cell']].value;w.close()
    elif loc['format']=='JSON':
        v=json.loads(raw,parse_float=Decimal)
        for key in loc['path']:v=v[key]
    else:raise ValueError('UNSUPPORTED_SOURCE_LOCATOR')
    if (None if v is None else str(v))!=record['value_original']:raise ValueError('LOCATOR_VALUE_MISMATCH')
    val,_=number(v)
    if loc['format']=='SEC_INLINE'and val is not None and record['attributes'].get('ix_sign')=='-':val=format(-Decimal(val),'f')
    if record.get('attributes',{}).get('footnote_pointer'):val=None
    if val!=record['value_decimal_or_null']:raise ValueError('DECIMAL_REPLAY_MISMATCH')
    return True


def backup_restore(ledger,destination,*,expected_current_hash=None):
    """Reuse Ledger.backup; source bytes and all eight logical tables checked."""
    before=file_sha(ledger.path)
    if expected_current_hash is not None and before!=expected_current_hash:raise ValueError('LOCAL_STATE_CAS_HASH_CHANGED')
    ledger.backup(destination)
    with sqlite3.connect(destination)as db:
        count=db.execute('select count(*) from events where kind=?',(OBS,)).fetchone()[0]
        if db.execute('pragma integrity_check').fetchone()!=('ok',):raise ValueError('RESTORE_INTEGRITY')
        for table in ('metadata','blobs','documents','captures','readings','runs','firsts','events'):
            a=ledger.db.execute('select * from '+table+' order by rowid').fetchall()
            b=db.execute('select * from '+table+' order by rowid').fetchall()
            if a!=b:raise ValueError('RESTORE_TABLE_MISMATCH')
    if file_sha(ledger.path)!=before:raise ValueError('ORIGINAL_CHANGED_DURING_BACKUP')
    return {'status':'PASS_LOCAL_8_TABLE_RESTORE','observations':count,'source_sha256':before,'backup_sha256':file_sha(destination),'not_remote_private_store':True}


def parse_awards(raw,meta):
    from bct.collectors.usaspending.collector import parse_response
    p=json.loads(raw,parse_float=Decimal);parse_response(raw,100);out=[]
    for i,r in enumerate(p['results']):
        ident=r['generated_internal_id'];date='UNKNOWN'
        out.append(observation(meta,ident,r.get('NAICS')or'UNKNOWN','award_amount',r.get('Award Amount'),'USD','MONEY',date,
          {'format':'JSON','path':['results',i,'Award Amount']},stock_or_flow='AWARD_CUMULATIVE_NOT_TRANSACTION_FLOW',
          attributes={'award_id':r['Award ID'],'award_type':r.get('Contract Award Type','UNKNOWN'),'PSC':r.get('PSC','UNKNOWN'),'agency':r.get('Awarding Agency','UNKNOWN'),'contract_start_not_observation_period':r.get('Start Date','UNKNOWN'),'contract_end_not_delivery':r.get('End Date','UNKNOWN'),'last_modified':r.get('Last Modified Date','UNKNOWN'),'physical_quantity':'UNKNOWN'},status='CONTRACT_AWARD_VALUE_NOT_PHYSICAL_DEMAND'))
    return out,p


def parse_transactions(raw,meta,award):
    p=json.loads(raw,parse_float=Decimal);out=[]
    for i,r in enumerate(p['results']):
        code=str(r.get('type',''))
        if code not in {'A','B','C','D'}:continue
        out.append(observation(meta,str(r['id']),award,'transaction_obligation',r.get('federal_action_obligation'),'USD','MONEY',r.get('action_date')or'UNKNOWN',
          {'format':'JSON','path':['results',i,'federal_action_obligation']},stock_or_flow='MODIFICATION_INCREMENT_NOT_AWARD_TOTAL',
          attributes={'award_id':award,'modification_number':r.get('modification_number'),'action_type':r.get('action_type'),'contract_type':code,'physical_quantity':'UNKNOWN'},status='CONTRACT_TRANSACTION_NOT_PHYSICAL_QUANTITY'))
    return out,p


def parse_ted(raw,meta):
    """Scalar totals only. Lot arrays without explicit identifiers stay held."""
    p=json.loads(raw,parse_float=Decimal);out=[];held=[]
    if p.get('timedOut'):raise ValueError('TED_TIMED_OUT_INCOMPLETE_PAGE')
    if not isinstance(p.get('notices'),list):raise ValueError('TED_RESPONSE_CONTRACT')
    for i,r in enumerate(p['notices']):
        ident=r.get('publication-number');pub=r.get('publication-date')
        if not isinstance(ident,str)or not isinstance(pub,str):held.append({'notice_index':i,'reason':'NOTICE_ID_DATE_NOT_SCALAR'});continue
        m=dict(meta,published_at=pub)
        v=r.get('total-value');currency=r.get('total-value-cur');path=['notices',i,'total-value']
        if isinstance(v,list)and len(v)==1:v=v[0];path.append(0)
        if isinstance(currency,list)and len(currency)==1:currency=currency[0]
        if isinstance(v,(str,int,float,Decimal))and isinstance(currency,str):
            out.append(observation(m,ident,str(r.get('classification-cpv','UNKNOWN')),'notice_total_value',v,currency,'MONEY',pub[:10],
              {'format':'JSON','path':path},geography='EU_TED_NOTICE_SCOPE',stock_or_flow='NOTICE_VALUE_NOT_ACTUAL_ORDER',status='NOTICE_STATUS_UNVERIFIED',
              attributes={'notice_type':r.get('notice-type'),'is_actual_order':False,'is_framework_or_option':'UNKNOWN','lot_relationship':'NOT_INFERRED_FROM_ARRAY_POSITION'}))
        for field in ['quantity-lot','quantity-unit-lot']:
            if field in r:held.append({'notice_id':ident,'field':field,'reason':'LOT_ID_UNIT_RELATION_REQUIRES_XML_NOT_ARRAY_ZIP'})
    for row in out:row['parser_version']='bct-q1-ted-exact-decimal-2'
    return out,p,held


def parse_tungsten(raw,meta,start,end):
    """USGS T1-3 only, explicit tungsten-content units, no macros/embedded code."""
    import openpyxl
    import calendar
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True);out=[];limitations=[]
    for name in ['T1','T2','T3']:
        if name not in w.sheetnames:limitations.append(name+'_MISSING');continue
        s=w[name];unit_text=str(s['A2'].value)
        if not re.search(r'metric tons, tungsten content',unit_text,re.I):
            limitations.append(name+'_UNIT_NOT_VERIFIED');continue
        if name=='T1':
            year=None
            for cells in s.iter_rows(min_row=5):
                label=str(cells[0].value or '').strip()
                if re.fullmatch(r'20\d{2}',label):year=label;continue
                month=next((i for i in range(1,13)if label==calendar.month_name[i]),None)
                if not year or not month:continue
                ym=year+'-%02d'%month
                if not start<=ym<=end:continue
                for j in range(1,6):
                    cell=cells[j];head=str(s.cell(4,j+1).value or 'UNKNOWN');cls='Concentrate'if j<=3 else'Intermediate products'
                    lineage='US_CENSUS_INTLTRADE'if'Imports'in head else'USGS_TUNGSTEN_SURVEY'
                    r=observation(meta,'US/'+cls,'tungsten/'+cls,head,cell.value,'metric tons tungsten content','PHYSICAL',ym,
                      {'format':'XLSX','sheet':name,'cell':cell.coordinate},stock_or_flow='STOCK_MONTH_END'if'Stocks'in head else'MONTHLY_FLOW',status='PRELIMINARY',lineage=lineage,
                      attributes={'unit_locator':name+'!A2','metric_locator':name+'!'+s.cell(4,j+1).coordinate,'is_total':False,'scope':'USGS table; actual underlying survey/trade lineage retained'})
                    if re.fullmatch(r'\(\d+\)',str(cell.value or '').strip()):r.update(value_decimal_or_null=None,value_normalized_or_null=None,quality_state='HELD',unknown_reason='FOOTNOTE_REFERENCE_NOT_NEGATIVE_QUANTITY');r['attributes']['footnote_pointer']=True
                    out.append(r)
        else:
            # The current reporting month is explicitly printed in columnA;
            # country rows beneath it belong only to this section, not prior totals.
            ym=None;columns=range(1,8 if name=='T2' else 7)
            for cells in s.iter_rows(min_row=4):
                label=str(cells[0].value or '').strip();m=re.fullmatch(r'(20\d{2}),\s*([A-Za-z]+)',label)
                if m:
                    mo=next((i for i in range(1,13)if m[2]==calendar.month_name[i]),None);ym=m[1]+'-%02d'%mo if mo else None;continue
                if not ym or not start<=ym<=end or not label or re.match(r'^\d|^Source:|^Note:',label):continue
                for j in columns:
                    if j>=len(cells):continue
                    cell=cells[j];head=s.cell(3,j+1).value
                    if not head:continue
                    r=observation(meta,label,'tungsten/'+str(head),'imports_for_consumption'if name=='T2'else'domestic_exports',cell.value,'metric tons tungsten content','PHYSICAL',ym,
                      {'format':'XLSX','sheet':name,'cell':cell.coordinate},stock_or_flow='MONTHLY_FLOW',status='PRELIMINARY',geography='US/from '+label if name=='T2'else'US/to '+label,lineage='US_CENSUS_INTLTRADE',
                      attributes={'unit_locator':name+'!A2','classification_locator':name+'!'+s.cell(3,j+1).coordinate,'is_country_total':label=='Total','is_product_total':str(head)=='Total','do_not_sum_parent_and_children':True})
                    if re.fullmatch(r'\(\d+\)',str(cell.value or '').strip()):r.update(value_decimal_or_null=None,value_normalized_or_null=None,quality_state='HELD',unknown_reason='FOOTNOTE_REFERENCE_NOT_NEGATIVE_QUANTITY');r['attributes']['footnote_pointer']=True
                    out.append(r)
    w.close();return out,limitations


def research_export(ledger):
    """Read-only numeric handoff; validation history wins over an early draft."""
    from bct.objective_lock import stamp_export
    latest={};validation={}
    for kind,payload in ledger.db.execute('select kind,payload from events where kind in (?,?) order by seq',(OBS,'QUANTITATIVE_VALIDATION')):
        event=json.loads(payload)
        if kind==OBS:latest[event['key']]=event
        else:
            for rid in event.get('record_ids',[]):validation[rid]=event
    out=[]
    for item in latest.values():
        r=dict(item['record']);v=validation.get(item['id'])
        if v:
            r.update(quality_state=v['quality_state'],validation_reason=v['reason'])
            if v['quality_state']=='HELD_OBSERVATION_PERIOD_UNKNOWN':
                r['attributes']={**r.get('attributes',{}),'preserved_draft_period':[r['period_start'],r['period_end']]}
                r.update(period_start='UNKNOWN',period_end='UNKNOWN')
        out.append({'id':item['id'],'record':r})
    return stamp_export({'plan_id':'BCT-V4-Q-1.0','step':'Q1','research_only':True,'observations':out,'candidates':[],
      'meaning':'Statistics and numeric source observations; not qualified supply, physical demand or forecast candidates.'})


def parse_hs_mapping(raw,meta,root_prefix):
    """Official HTS10 metadata, not trade quantities or interchangeable Schedule B.

    Prefixes must be present headings in this exact official export. No inferred
    commodity code, no kg/W-kg conversion, no unit inherited from a parent total.
    """
    if not re.fullmatch(r'\d{4}|\d{6}|\d{8}|\d{10}',root_prefix):raise ValueError('HS_ROOT_CONTRACT')
    rows=json.loads(raw)
    if not isinstance(rows,list):raise ValueError('HS_TABLE_CONTRACT')
    if not any(str(r.get('htsno','')).replace('.','')==root_prefix for r in rows):raise ValueError('HS_ROOT_NOT_IN_OFFICIAL_TABLE')
    out=[]
    for i,r in enumerate(rows):
        code=str(r.get('htsno','')).replace('.','')
        if not re.fullmatch(r'\d{10}',code)or not code.startswith(root_prefix):continue
        units=r.get('units',[])
        if not 1<=len(units)<=2 or any(not isinstance(u,str)or not u.strip()for u in units):raise ValueError('HS_LEAF_UNITS_UNKNOWN')
        out.append({'import_HTS10':code,'display_code':r['htsno'],'description':r['description'],
                    'units_original':units,'primary_quantity_unit':units[0],
                    'secondary_quantity_unit':units[1]if len(units)>1 else None,
                    'source_url':meta['url'],'source_hash':meta['sha256'],
                    'source_locator':{'format':'JSON','path':[i]},'root_verified':root_prefix,
                    'version':'CURRENT_RETRIEVED_HASH_NOT_HISTORICAL_CODEBOOK',
                    'trade_quantities':'NOT_FETCHED','not_qualified_supply':True})
    return out


def quantity_role(record):
    """Descriptive reporting buckets; never a TARGET or shortage inference."""
    if record['value_decimal_or_null']is None:return 'WITHHELD_OR_UNKNOWN'
    kind=record['measurement_type'];source=record['source_id'];metric=record['metric']
    if kind=='MONEY':return 'MONEY_ONLY_NOT_PHYSICAL_DEMAND'
    if kind=='INDEX':return 'CAPACITY_INDEX'if metric=='industrial_capacity'else'PRODUCTION_INDEX'
    if kind=='RATIO':return 'UTILIZATION_RATIO'
    if source=='eia'and metric=='nameplate_capacity':
        return {'Operating':'OPERATING_NAMEPLATE_NOT_QUALIFIED','Planned':'PLANNED_NAMEPLATE_NOT_AVAILABLE',
                'Retired':'RETIRED_NAMEPLATE','Canceled or Postponed':'CANCELED_OR_POSTPONED_NAMEPLATE'}.get(record['attributes'].get('sheet'),'CAPACITY_STATUS_UNKNOWN')
    if source=='g17'and metric=='vehicle_assemblies':return 'REPORTED_PRODUCTION_ANNUALIZED_RATE'
    if source=='usgs':
        if metric=='Consumption':return 'REPORTED_CONSUMPTION_NOT_CUSTOMER_ORDER'
        if metric.startswith('Net production'):return 'REPORTED_PRODUCTION_FLOW'
        if metric.startswith('Stocks'):return 'REPORTED_INVENTORY_STOCK'
        if 'import'in metric.lower()or'export'in metric.lower():return 'REPORTED_TRADE_FLOW_NOT_QUALIFIED_SUPPLY'
    return 'PHYSICAL_ROLE_UNKNOWN'
